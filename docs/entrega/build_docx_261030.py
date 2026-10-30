#!/usr/bin/env python3
"""Gera docs/entrega/ProjetoG2_261030.docx (entrega da aula de 30/10/2026: model / persistência).

Entradas: capturas em docs/entrega/telas/*.png (`npm run e2e:shots`), diagramas em diagramas/*.png (mermaid)
e trechos de código lidos direto de apps/api/src (marcadores de início/fim). Uso: python3 docs/entrega/build_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).parent
TELAS = BASE / "telas"
DIAG = BASE / "diagramas"
OUT = BASE / "ProjetoG2_261030.docx"
API = BASE.parent.parent / "apps" / "api" / "src"

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x6B, 0x82)

TEAM = [
    ("Gustavo Peixoto Siriacov", "202411977", "responsável pelo envio no LEX"),
    ("Tales Ivaniski Galina", "202510926", ""),
    ("Lorenzo Ruggini Lindemann", "202511046", ""),
    ("Eduardo Cassel da Costa", "202513138", ""),
]


def img(path: Path) -> Path:
    if not path.exists():
        raise SystemExit(f"Imagem ausente: {path} (rode `npm run e2e:shots` e o mermaid antes)")
    return path


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def para(doc, text="", bold=False, size=None, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, label, text=""):
    p = doc.add_paragraph(style="List Bullet")
    if text:
        r = p.add_run(label)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(label)
    p.paragraph_format.space_after = Pt(2)
    return p


def caption(doc, text):
    p = para(doc, text, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, space_after=10)
    return p


def shots_row(doc, items, width_cm):
    """Linha de capturas lado a lado: items = [(arquivo, legenda), ...]."""
    table = doc.add_table(rows=2, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (name, cap) in enumerate(items):
        cell = table.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(img(TELAS / name)), width=Cm(width_cm))
        c = table.cell(1, i).paragraphs[0]
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap)
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = MUTED
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, "1B3A5F")
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t




def snip(rel: str, start: str, end: str | None = None, extra: int = 0) -> str:
    """Trecho de apps/api/src/<rel>: da 1ª linha com `start` até a 1ª linha seguinte com `end` (inclusive)."""
    lines = (API / rel).read_text(encoding="utf-8").splitlines()
    i = next((k for k, l in enumerate(lines) if start in l), None)
    if i is None:
        raise SystemExit(f"Marcador não encontrado em {rel}: {start!r}")
    j = i
    if end:
        j = next((k for k in range(i, len(lines)) if end in lines[k]), None)
        if j is None:
            raise SystemExit(f"Fim não encontrado em {rel}: {end!r}")
    block = lines[i : j + 1 + extra]
    cut = min((len(l) - len(l.lstrip()) for l in block if l.strip()), default=0)
    return "\n".join(l[cut:] for l in block)


def code(doc, label: str, text: str, size=8):
    """Bloco de código: título com o arquivo de origem e texto monoespaçado sobre fundo cinza."""
    para(doc, label, bold=True, size=9, color=MUTED, space_after=1)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, "F3F5F8")
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for n, line in enumerate(text.split("\n")):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(size)
        if n < text.count("\n"):
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pending(doc, text):
    """Campo a ser preenchido pelo usuário, destacado em amarelo."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW
    return p


def fig(doc, name, cap, width_cm=None, height_cm=None, folder=TELAS):
    kw = {"width": Cm(width_cm)} if width_cm else {"height": Cm(height_cm)}
    doc.add_picture(str(img(folder / name)), **kw)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, cap)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY

# ------------------------------------------------------------------ identificação
para(doc, "Projeto G2: Organizador de Estudos com IA Integrada", bold=True, size=22, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, "ProjetoG2_261030 · Entrega da aula de 30/10/2026", size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, "Desenvolvimento de Dispositivos Móveis · Universidade La Salle", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ------------------------------------------------------------------ 1. integrantes
doc.add_heading("1. Integrantes que colaboraram na aula de hoje (30/10)", level=1)
para(doc, "Aplicativo: EduMind. Integrantes da equipe (conforme a entrega de 23/10):", space_after=4)
table(
    doc,
    ["Integrante", "Matrícula", "Colaborou hoje?"],
    [[n, m, "Sim"] for n, m, _ in TEAM],
    widths=[7.5, 3, 6],
)
para(doc, "Todos os integrantes colaboraram na aula de 30/10.", space_after=6)

# ------------------------------------------------------------------ 2. model
doc.add_heading("2. O que é o model do EduMind", level=1)
para(
    doc,
    "O model é a camada que guarda e lê os dados. No EduMind ele fica na API (apps/api/src), em SQLite, e o app web "
    "só conversa com ela por HTTP: nenhuma tela acessa o banco diretamente. Os dados ficam no servidor, ligados à conta "
    "do usuário, e por isso aparecem em qualquer dispositivo com o mesmo login.",
)
bullet(doc, "Estrutura (db/schema.sql): ", "12 tabelas com chaves estrangeiras, restrições (unique, check) e índices. É aplicada pelo migrador no primeiro uso do banco (pragma user_version).")
bullet(doc, "Acesso (db/db.ts): ", "quatro atalhos de consulta com parâmetros (one, all, run, insert). Nenhum valor é concatenado no SQL, o que evita injeção de SQL.")
bullet(doc, "Repositórios e rotas (modules/*): ", "cada módulo tem as operações da sua tabela. Antes de ler ou alterar, a API confirma que o registro pertence ao usuário logado.")
bullet(doc, "Integridade: ", "apagar uma disciplina apaga em cascata seus tópicos, metas, aulas, trabalhos e sessões de estudo; o log de acessos sobrevive à exclusão da conta (guarda cópia do nome e do e-mail).")

# DER em página paisagem
new = doc.add_section()
new.orientation = WD_ORIENT.LANDSCAPE
new.page_width, new.page_height = Cm(29.7), Cm(21)
new.left_margin = new.right_margin = Cm(1.5)
new.top_margin = new.bottom_margin = Cm(1.5)
doc.add_heading("3. Diagrama do modelo de dados (DER)", level=1)
fig(doc, "der-modelo.png", "Figura 1 — Tabelas persistidas e seus relacionamentos (cada usuário só enxerga os próprios dados).", width_cm=26.5, folder=DIAG)
back = doc.add_section()
back.orientation = WD_ORIENT.PORTRAIT
back.page_width, back.page_height = Cm(21), Cm(29.7)
back.left_margin = back.right_margin = Cm(2.2)
back.top_margin = back.bottom_margin = Cm(2)

para(doc, "Resumo das tabelas:", bold=True, space_after=4)
table(
    doc,
    ["Tabela", "O que guarda"],
    [
        ["users", "conta: e-mail, nome, hash da senha (scrypt), papel, ativo/bloqueado, meta semanal"],
        ["sessions", "uma sessão por dispositivo (hash do token, validade, encerramento)"],
        ["access_logs", "log de acesso: nome, data, hora de login e de logout"],
        ["subjects", "disciplinas do usuário (nome único por usuário, professor, sala, cor)"],
        ["topics", "tópicos de cada disciplina com status (pendente, estudando, concluído) e prazo"],
        ["daily_goals", "metas do dia por disciplina, com marcação de concluída"],
        ["study_sessions", "minutos estudados por dia e disciplina (base da tela Progresso)"],
        ["chat_threads / chat_messages", "conversas com o Professor Tutor IA e suas mensagens"],
        ["lessons / assignments", "aulas e trabalhos lidos da faculdade pelo scraper"],
        ["sync_runs", "histórico das sincronizações (status, início, fim, estatísticas)"],
    ],
    widths=[4.5, 12.5],
)

# ------------------------------------------------------------------ 3.x estrutura detalhada
import re

DESC = {
    "users": ("Contas de usuário. Cada estudante e cada administrador é uma linha.", {
        "id": "identificador gerado automaticamente", "email": "login; único, sem diferenciar maiúsculas",
        "name": "nome exibido no app e no log", "password_hash": "senha protegida com scrypt (a senha em si nunca é guardada)",
        "level": "nível de ensino: medio, superior ou concurso", "role": "papel: user ou admin",
        "is_active": "1 = ativa, 0 = bloqueada pelo administrador", "weekly_goal_minutes": "meta de estudo da semana, em minutos",
        "created_at": "data e hora do cadastro", "last_login_at": "data e hora do último login"}),
    "sessions": ("Uma linha por dispositivo logado. O servidor consegue encerrar qualquer sessão.", {
        "id": "identificador", "user_id": "dono da sessão (apagada junto com a conta)", "token_hash": "hash sha256 do token do cookie; o token em si não é guardado",
        "device": "texto como \"Chrome · Android\"", "user_agent": "identificação completa do navegador", "ip": "endereço do acesso",
        "created_at": "início da sessão", "last_seen_at": "última atividade", "expires_at": "validade (7 dias por padrão)",
        "revoked_at": "quando foi encerrada (vazio = ativa)", "revoke_reason": "motivo: manual, expired, revoked ou admin"}),
    "access_logs": ("Histórico de acessos: nome, data, hora de login e hora de logout. Sobrevive à exclusão da conta.", {
        "id": "identificador", "user_id": "usuário (vira vazio se a conta for excluída)", "user_name": "cópia do nome no momento do login",
        "user_email": "cópia do e-mail no momento do login", "session_id": "sessão que originou o registro", "method": "password ou google",
        "device": "dispositivo usado", "ip": "endereço do acesso", "login_at": "data e hora do login",
        "logout_at": "data e hora do logout (vazio = ainda em andamento)", "logout_reason": "manual, expired, revoked ou admin"}),
    "subjects": ("Disciplinas do estudante, criadas à mão ou importadas pelo scraper.", {
        "id": "identificador", "user_id": "dono da disciplina", "name": "nome; único por usuário, sem diferenciar maiúsculas",
        "professor": "nome do professor (opcional)", "room": "sala (opcional)", "color": "cor em hexadecimal (#RRGGBB)",
        "source": "origem: manual ou scraper", "external_id": "código na faculdade, usado para não duplicar importações", "created_at": "data e hora de criação"}),
    "topics": ("Tópicos (assuntos) de estudo de cada disciplina.", {
        "id": "identificador", "subject_id": "disciplina do tópico (apagado junto com ela)", "title": "título do tópico", "notes": "anotações (opcional)",
        "position": "ordem de exibição dentro da disciplina", "status": "pendente, estudando ou concluido", "due_date": "prazo (AAAA-MM-DD, opcional)",
        "done_at": "quando foi concluído (preenchido e limpo automaticamente)"}),
    "daily_goals": ("Metas do dia, ligadas a uma disciplina.", {
        "id": "identificador", "user_id": "dono da meta", "subject_id": "disciplina da meta", "goal_date": "dia da meta (AAAA-MM-DD)",
        "description": "o que estudar", "target_minutes": "tempo previsto em minutos (opcional, maior que zero)", "done": "1 = concluída",
        "done_at": "quando foi concluída"}),
    "study_sessions": ("Tempo realmente estudado; é a base dos gráficos da tela Progresso.", {
        "id": "identificador", "user_id": "estudante", "subject_id": "disciplina estudada", "started_at": "data e hora de início",
        "day": "dia local do aluno, usado para somar por dia", "minutes": "duração em minutos (maior que zero)"}),
    "chat_threads": ("Conversas com o Professor Tutor IA.", {
        "id": "identificador", "user_id": "dono da conversa", "subject_id": "disciplina de contexto (vazio = conversa geral)",
        "title": "título gerado a partir da primeira mensagem", "created_at": "criação", "updated_at": "última mensagem"}),
    "chat_messages": ("Mensagens de cada conversa.", {
        "id": "identificador", "thread_id": "conversa (apagada junto com ela)", "role": "user (aluno) ou assistant (tutor)",
        "content": "texto da mensagem", "created_at": "data e hora"}),
    "lessons": ("Aulas lidas da faculdade pelo scraper.", {
        "id": "identificador", "user_id": "estudante", "subject_id": "disciplina da aula", "external_id": "código da aula na faculdade; único por usuário",
        "date": "dia da aula", "start_time": "hora de início (HH:mm)", "end_time": "hora de término (HH:mm)", "topic": "assunto (opcional)",
        "room": "sala (opcional)", "updated_at": "última atualização pela importação"}),
    "assignments": ("Trabalhos e provas lidos da faculdade pelo scraper.", {
        "id": "identificador", "user_id": "estudante", "subject_id": "disciplina do trabalho", "external_id": "código na faculdade; único por usuário",
        "title": "título", "description": "descrição (opcional)", "due_date": "data de entrega", "status": "pendente, entregue ou corrigido",
        "grade": "nota (opcional)", "updated_at": "última atualização pela importação"}),
    "sync_runs": ("Registro de cada sincronização com a faculdade.", {
        "id": "identificador", "user_id": "quem sincronizou", "source": "mock (portal de demonstração) ou portal", "status": "running, ok ou error",
        "started_at": "início", "finished_at": "fim", "stats_json": "resumo (quantos itens criados/atualizados)", "error": "mensagem de erro, se houver"}),
}


def schema_columns() -> dict:
    """Colunas de cada tabela, lidas de db/schema.sql: {tabela: [(coluna, tipo, regras)]}."""
    txt = (API / "db" / "schema.sql").read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"create table (\w+) \((.*?)\n\);", txt, re.S):
        cols = []
        for line in m.group(2).split("\n"):
            line = re.sub(r"--.*", "", line).strip().rstrip(",")
            mm = re.match(r"(\w+)\s+(integer|text|real)\b(.*)", line)
            if mm:
                rules = mm.group(3).strip()
                rules = re.sub(r"\s+", " ", rules)
                cols.append((mm.group(1), mm.group(2), rules))
        out[m.group(1)] = cols
    return out


doc.add_heading("3.1 Como o banco é criado e usado", level=2)
para(doc, "Do primeiro uso até a resposta que o app recebe, o caminho dos dados é este:", space_after=4)
for label, text in [
    ("1. Abertura: ", "ao iniciar, a API abre o arquivo SQLite (data/edumind.db), liga a verificação de chaves estrangeiras e o modo WAL (leituras e escritas simultâneas)."),
    ("2. Migração: ", "o banco guarda um número de versão (pragma user_version). Se estiver abaixo da versão atual, o arquivo schema.sql é executado dentro de uma transação, criando todas as tabelas e índices. Se já estiver atualizado, nada acontece."),
    ("3. Requisição: ", "o app web envia a ação (por exemplo, criar disciplina) por HTTP. A API confirma o login pelo cookie de sessão."),
    ("4. Validação: ", "os dados enviados são conferidos por um esquema (tipos, tamanhos, formatos de data). Dado inválido é recusado antes de chegar ao banco."),
    ("5. Operação: ", "o repositório/rota executa o SQL (insert, select, update ou delete) com parâmetros e sempre filtrando pelo usuário logado, de modo que ninguém acessa dados de outra conta."),
    ("6. Resposta: ", "a linha do banco (nomes como weekly_goal_minutes) é convertida para o formato do app (weeklyGoalMinutes) e devolvida em JSON."),
]:
    bullet(doc, label, text)
para(doc, "Convenções de armazenamento:", bold=True, space_after=2)
bullet(doc, "Data e hora: ", "texto ISO-8601 em UTC (ex.: 2026-10-30T14:05:00.000Z), sempre gravado pela aplicação.")
bullet(doc, "Dia de calendário: ", "AAAA-MM-DD (data local do aluno); horário de aula: HH:mm.")
bullet(doc, "Verdadeiro/falso: ", "inteiro 0 ou 1 (SQLite não tem tipo booleano).")
bullet(doc, "Tipos usados: ", "integer (números), text (textos e datas) e real (nota).")
bullet(doc, "Restrições no próprio banco: ", "not null, unique, check (valores permitidos), chaves estrangeiras com on delete cascade ou set null, e índices para consultas rápidas.")

doc.add_heading("3.2 Tabelas e colunas", level=2)
para(doc, "Para cada tabela: finalidade e colunas (tipo, regras do banco e significado). PK = chave primária; FK = chave estrangeira.", color=MUTED)
cols = schema_columns()
for tname, (purpose, cdesc) in DESC.items():
    if tname not in cols:
        raise SystemExit(f"Tabela ausente no schema.sql: {tname}")
    para(doc, tname, bold=True, size=12, color=NAVY, space_after=1)
    para(doc, purpose, space_after=3)
    rows = []
    for c, t, rules in cols[tname]:
        rule = rules
        rule = re.sub(r"primary key autoincrement", "PK", rule)
        rule = re.sub(r"references (\w+)\((\w+)\)( on delete (cascade|set null))?", lambda m: f"FK → {m.group(1)}" + (f" ({m.group(4)})" if m.group(4) else ""), rule)
        rule = re.sub(r"check \(.*", "restrito a valores válidos", rule) if "check" in rule else rule
        rows.append([c, t, rule.replace("not null", "obrigatório").replace("default", "padrão").replace("unique", "único").replace("collate nocase", "sem diferenciar caixa").strip(), cdesc.get(c, "")])
    table(doc, ["Coluna", "Tipo", "Regras", "Significado"], rows, widths=[3.2, 1.6, 4.4, 7.8])

# ------------------------------------------------------------------ 4. telas
doc.add_page_break()
doc.add_heading("4. Telas envolvidas na persistência", level=1)
para(doc, "Cada tela abaixo grava ou lê dados no banco. A tabela relaciona a ação do usuário com a operação feita no model.", color=MUTED)
table(
    doc,
    ["Tela", "Ação do usuário", "Operação no banco"],
    [
        ["Cadastro", "cria a conta", "INSERT em users; abre sessão e log (INSERT em sessions e access_logs)"],
        ["Login", "entra no app", "SELECT em users; INSERT em sessions e access_logs; UPDATE de users.last_login_at"],
        ["Disciplinas", "cria, edita e exclui disciplina", "INSERT / SELECT / UPDATE / DELETE em subjects"],
        ["Detalhe da disciplina", "adiciona tópico e muda o status", "INSERT / SELECT / UPDATE / DELETE em topics"],
        ["Início (metas do dia)", "cria e conclui metas", "INSERT / SELECT / UPDATE / DELETE em daily_goals"],
        ["Progresso", "registra tempo e vê a semana", "INSERT em study_sessions; SELECT com soma por dia e disciplina"],
        ["Perfil e acessos", "edita o perfil e vê os dispositivos", "UPDATE em users; SELECT em sessions e access_logs; encerra sessão (UPDATE)"],
        ["Tutor IA", "conversa com o tutor", "INSERT em chat_threads e chat_messages; DELETE da conversa"],
        ["Aulas", "sincroniza aulas e trabalhos", "INSERT ou UPDATE em lessons, assignments e subjects; registro em sync_runs"],
        ["Admin", "gerencia usuários e vê o log", "SELECT em access_logs e sessions; UPDATE e DELETE em users"],
    ],
    widths=[3.6, 5.0, 8.4],
)

doc.add_heading("Cadastro e login (users, sessions, access_logs)", level=2)
shots_row(doc, [("03-cadastro.png", "Cadastro"), ("01-login.png", "Login"), ("10-perfil-acessos.png", "Perfil e acessos")], width_cm=5.0)
doc.add_heading("Disciplinas, detalhe e metas (subjects, topics, daily_goals)", level=2)
shots_row(doc, [("06-disciplina-form.png", "Nova disciplina"), ("07-disciplinas.png", "Disciplinas"), ("08-disciplina-detalhe.png", "Detalhe (tópicos)")], width_cm=5.0)
shots_row(doc, [("05-dashboard-vazio.png", "Início vazio"), ("11-dashboard.png", "Início com dados (desktop)")], width_cm=7.0)
doc.add_heading("Progresso e tutor (study_sessions, chat_*)", level=2)
shots_row(doc, [("09-progresso.png", "Progresso"), ("14-tutor.png", "Tutor IA (desktop)")], width_cm=6.5)
doc.add_heading("Aulas sincronizadas (lessons, assignments)", level=2)
fig(doc, "13-aulas.png", "Figura 2 — Aulas e trabalhos importados pelo scraper (desktop).", width_cm=12.5)
doc.add_heading("Administração (users, sessions, access_logs)", level=2)
shots_row(doc, [("16-admin-log-acessos.png", "Log de acessos"), ("17-admin-usuarios.png", "Usuários")], width_cm=7.8)
shots_row(doc, [("15-admin-visao-geral.png", "Visão geral"), ("18-admin-sessoes.png", "Sessões ativas")], width_cm=7.8)

# ------------------------------------------------------------------ 5. código
doc.add_page_break()
doc.add_heading("5. Código do model: operações com os dados", level=1)
para(doc, "Trechos reais da API (apps/api/src), reduzidos ao essencial. Todas as consultas usam parâmetros (?).", color=MUTED)

doc.add_heading("5.1 Conexão, migração e atalhos de consulta", level=2)
code(doc, "db/db.ts — migração do esquema", snip("db/db.ts", "export function migrate", "    })();", 2))
code(doc, "db/db.ts — atalhos usados por todos os módulos", snip("db/db.ts", "export const one", "Number(run(db, sql"))
code(doc, "db/schema.sql — exemplo de tabela (restrições e chave estrangeira)", snip("db/schema.sql", "create table subjects", ");", 1))

doc.add_heading("5.2 INSERT (create)", level=2)
code(doc, "modules/auth/users.ts — cria o usuário (tela Cadastro)", snip("modules/auth/users.ts", "export function insertUser", "return findUserById", 1))
code(doc, "modules/subjects/routes.ts — cria a disciplina (tela Disciplinas)", snip("modules/subjects/routes.ts", "const id = insert(", "ctx.now(),", 1))
code(doc, "modules/topics/routes.ts — cria o tópico na próxima posição", snip("modules/topics/routes.ts", "const next = one", "subject.id, input.title", 1))
code(doc, "modules/study/routes.ts — registra uma sessão de estudo (tela Progresso)", snip("modules/study/routes.ts", "const id = insert(", "input.minutes,", 1))

doc.add_heading("5.3 SELECT (read)", level=2)
code(doc, "modules/subjects/repo.ts — disciplina com contagem de tópicos, só do dono", snip("modules/subjects/repo.ts", "export const SUBJECT_SQL", "getSubject = (db", 1))
code(doc, "modules/goals/routes.ts — metas de um dia", snip("modules/goals/routes.ts", "const rows = all<GoalRow>(", "req.auth!.user.id, date,", 1))
code(doc, "modules/access/repo.ts — log de acessos com filtros e paginação", snip("modules/access/repo.ts", "export function listAccessLogs", "return { items: rows.map(toAccessLog), total };", 1))

doc.add_heading("5.4 UPDATE", level=2)
code(doc, "modules/topics/routes.ts — muda o status e mantém done_at coerente", snip("modules/topics/routes.ts", "const status = patch.status", "current.id,", 1))
code(doc, "modules/auth/routes.ts — atualiza o perfil (tela Perfil)", snip("modules/auth/routes.ts", "run(", "patch.name ?? null", 1))

doc.add_heading("5.5 DELETE", level=2)
code(doc, "modules/subjects/routes.ts — exclui a disciplina (apaga em cascata)", snip("modules/subjects/routes.ts", "// Tópicos, metas", "res.status(204)", 1))
code(doc, "modules/admin/routes.ts — exclui usuário (transação)", snip("modules/admin/routes.ts", "access.endAllForUser(target.id, 'admin'); // fecha", "run(db, 'delete from users"))

doc.add_heading("5.6 Transações e importação", level=2)
code(doc, "modules/access/service.ts — login: sessão + log + último acesso, tudo ou nada", snip("modules/access/service.ts", "const sessionId = db.transaction", "return id;", 1))
code(doc, "modules/access/service.ts — logout: encerra a sessão e fecha o log", snip("modules/access/service.ts", "function endSession", "})();", 1))
code(doc, "modules/sync/import.ts — importação idempotente de aulas (insere ou atualiza)", snip("modules/sync/import.ts", "for (const l of course.lessons)", "stats.updated++;", 2))

# ------------------------------------------------------------------ 6. pendências
doc.add_heading("6. Observação", level=1)
para(
    doc,
    "As telas e os trechos refletem o estado do projeto em 30/10/2026. O model continua sendo desenvolvido; "
    "os trechos de código deste documento são gerados a partir dos arquivos da API a cada execução do script de build.",
)

doc.save(OUT)
print(f"Gerado: {OUT}")
