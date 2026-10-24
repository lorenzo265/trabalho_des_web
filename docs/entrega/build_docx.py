#!/usr/bin/env python3
"""Gera docs/entrega/ProjetoG2_261023.docx (entrega da aula de 23/10/2026).

Entradas: diagramas em docs/entrega/diagramas/*.png e capturas em docs/entrega/telas/*.png
(gerados por `npm run e2e:shots` e pelo mermaid). Uso: python3 docs/entrega/build_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).parent
TELAS = BASE / "telas"
DIAG = BASE / "diagramas"
OUT = BASE / "ProjetoG2_261023.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x6B, 0x82)

TEAM = [
    ("Gustavo Peixoto Siriacov", "202411977", "responsável pelo envio no LEX"),
    ("Tales Ivaniski Galina", "202510926", ""),
    ("Lorenzo Ruggini Lindemann", "202511046", ""),
    ("Eduardo Cassel da Costa", "202513138", ""),
]


def img(path: Path) -> Path:
    if not path.exists():
        raise SystemExit(f"Imagem ausente: {path} (rode `npm run e2e:shots` e o mermaid antes)")
    return path


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def para(doc, text="", bold=False, size=None, color=None, align=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, label, text=""):
    p = doc.add_paragraph(style="List Bullet")
    if text:
        r = p.add_run(label)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(label)
    p.paragraph_format.space_after = Pt(2)
    return p


def caption(doc, text):
    p = para(doc, text, size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, space_after=10)
    return p


def shots_row(doc, items, width_cm):
    """Linha de capturas lado a lado: items = [(arquivo, legenda), ...]."""
    table = doc.add_table(rows=2, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (name, cap) in enumerate(items):
        cell = table.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(img(TELAS / name)), width=Cm(width_cm))
        c = table.cell(1, i).paragraphs[0]
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap)
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = MUTED
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, "1B3A5F")
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY

# ------------------------------------------------------------------ identificação
para(doc, "Projeto G2: Organizador de Estudos com IA Integrada", bold=True, size=22, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, "ProjetoG2_261023 · Entrega da aula de 23/10/2026", size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(doc, "Desenvolvimento de Dispositivos Móveis · Universidade La Salle", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

doc.add_heading("1. Equipe", level=1)
para(doc, "Nome da equipe / do aplicativo: EduMind", bold=True, space_after=6)
table(doc, ["Integrante", "Matrícula", "Observação"], [list(m) for m in TEAM], widths=[7.5, 3, 6])

doc.add_heading("2. Categoria do aplicativo", level=1)
para(
    doc,
    "O EduMind é um web app responsivo (mobile-first): roda na nuvem/servidor e é aberto pelo navegador do dispositivo, "
    "podendo ser fixado na tela inicial como um ícone de aplicativo. Os dados ficam no servidor, ligados à conta do usuário "
    "(e não ao aparelho), o que permite acessar de qualquer dispositivo com o mesmo e-mail e senha.",
)
bullet(doc, "Frontend: ", "React + Vite (TypeScript), layout responsivo do celular ao desktop.")
bullet(doc, "Backend: ", "Node + Express, banco SQLite, sessão por cookie.")
bullet(doc, "Aulas da faculdade: ", "scraper com Playwright (somente leitura), com um portal de demonstração para desenvolvimento.")
doc.add_picture(str(img(DIAG / "arquitetura.png")), width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption(doc, "Figura 1 — Divisão do projeto em três frentes: telas, acessos e aulas (scrapers).")

# ------------------------------------------------------------------ telas de acesso
doc.add_page_break()
doc.add_heading("3. Telas de acesso do usuário", level=1)
para(doc, "Interface minimalista: cada tela tem uma tarefa e poucos campos, com ícones padrão e ajuda contextual (botão “?”).", color=MUTED)

doc.add_heading("Tela 1 — Autenticação (Login)", level=2)
bullet(doc, "Descrição: ", "porta de entrada do app; identifica o estudante e isola os dados de cada conta.")
bullet(doc, "Dados exibidos: ", "logo e nome do aplicativo (EduMind), mensagem de boas-vindas, campos e botões de ação.")
bullet(doc, "Dados informados pelo usuário: ", "e-mail e senha. O botão “Entrar com Google” aparece desabilitado (“em breve”) nesta versão.")
bullet(doc, "Validações: ", "e-mail válido e senha obrigatória no próprio formulário; erro genérico quando as credenciais não conferem (não revela qual campo está errado) e limite de tentativas seguidas.")
shots_row(
    doc,
    [("01-login.png", "Login"), ("02-login-erro.png", "Erro de credenciais"), ("03-cadastro.png", "Cadastro com validação")],
    width_cm=5.0,
)

doc.add_heading("Cadastro e primeiros passos", level=2)
bullet(doc, "Cadastro: ", "nome, e-mail e senha (mínimo de 8 caracteres). A conta é criada e o usuário já entra no app.")
bullet(doc, "Apresentação inicial: ", "3 cartões exibidos uma vez por usuário (e que podem ser revistos no Perfil), seguindo a dica de aula sobre telas que precisam de ajuda/treinamento.")
shots_row(doc, [("04-tour.png", "Apresentação inicial"), ("05-dashboard-vazio.png", "Início após o cadastro")], width_cm=5.0)

doc.add_heading("Perfil: dispositivos conectados e meus acessos", level=2)
bullet(doc, "Dispositivos conectados: ", "lista cada aparelho com sessão aberta (ex.: Chrome · Android) e permite desconectar os outros. É o que torna o acesso independente do dispositivo.")
bullet(doc, "Meus acessos: ", "histórico do próprio usuário com data, hora de login, hora de logout, dispositivo e situação (em andamento, saiu do app, sessão expirada...).")
shots_row(doc, [("10-perfil-acessos.png", "Perfil, dispositivos e meus acessos"), ("11-dashboard.png", "Início (conta demo, desktop)")], width_cm=8.0)

# ------------------------------------------------------------------ modelo de dados
doc.add_page_break()
doc.add_heading("4. Modelo de dados do acesso e do log", level=1)
para(
    doc,
    "O acesso do usuário ao app é independente do dispositivo porque a conta (users) é única e cada dispositivo ganha a sua "
    "própria sessão (sessions). O log de acesso (access_logs) registra uma linha por login, com o nome do usuário, a data e "
    "hora do login e a hora do logout.",
)
doc.add_picture(str(img(DIAG / "der-acesso.png")), height=Cm(12.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption(doc, "Figura 2 — Diagrama entidade-relacionamento: usuários, sessões (uma por dispositivo) e log de acesso.")

table(
    doc,
    ["Tabela", "Para que serve", "Pontos de atenção"],
    [
        ["users", "Conta do estudante ou administrador (e-mail, nome, papel, ativo/bloqueado).", "Senha guardada só como hash (scrypt). O papel define o acesso à área do administrador."],
        ["sessions", "Um acesso aberto em um dispositivo (cookie de sessão, dispositivo, IP, validade de 7 dias).", "Só o hash do token é guardado. Encerrar a sessão (logout, desconectar aparelho, bloqueio) preenche revoked_at."],
        ["access_logs", "Histórico de acessos: nome, data e hora do login, hora do logout, dispositivo e motivo.", "logout_at nulo = em andamento. Nome e e-mail são cópias: o log sobrevive à exclusão da conta."],
    ],
    widths=[2.6, 7.2, 7.2],
)
para(doc, "Regras do log", bold=True, space_after=2)
bullet(doc, "Todo login bem-sucedido cria 1 sessão e 1 linha de log na mesma transação.")
bullet(doc, "O logout é gravado uma única vez: saiu do app, sessão expirada, dispositivo desconectado ou encerrada pelo administrador.")
bullet(doc, "Tentativas de login que falham não geram log (apenas contam no limite de tentativas).")
bullet(doc, "Se o navegador for fechado sem sair, o log fica “em andamento” até a sessão expirar.")

# ------------------------------------------------------------------ fluxograma
doc.add_page_break()
doc.add_heading("5. Fluxograma de acesso e registro do log", level=1)
doc.add_picture(str(img(DIAG / "fluxo-acesso-log.png")), height=Cm(21))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption(doc, "Figura 3 — Do login ao logout: criação da sessão do dispositivo e registro de data e hora no log.")

# ------------------------------------------------------------------ admin
doc.add_page_break()
doc.add_heading("6. Área do administrador", level=1)
para(
    doc,
    "Decisão: a área do administrador é incluída no mesmo app web (rota /admin), visível apenas para usuários com papel "
    "administrador. O layout é pensado primeiro para desktop (abas, tabelas e filtros) e continua utilizável no celular, "
    "onde as tabelas viram listas de cartões.",
)
bullet(doc, "Visão geral: ", "usuários, sessões ativas, logins de hoje e gráfico dos últimos 14 dias.")
bullet(doc, "Usuários: ", "busca, mudança de papel, bloqueio/ativação e exclusão (com salvaguardas: o admin não bloqueia nem exclui a si mesmo, nem o último administrador).")
bullet(doc, "Log de acessos: ", "nome, data, hora de login, hora de logout, duração, dispositivo e método; filtros por usuário e período; exportação em CSV.")
bullet(doc, "Sessões ativas: ", "dispositivos conectados agora, com opção de encerrar à força.")
doc.add_picture(str(img(TELAS / "16-admin-log-acessos.png")), width=Cm(12.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption(doc, "Figura 4 — Log de acessos na área do administrador.")
doc.add_picture(str(img(TELAS / "15-admin-visao-geral.png")), width=Cm(12.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption(doc, "Figura 5 — Visão geral da administração.")

# ------------------------------------------------------------------ demais telas
doc.add_page_break()
doc.add_heading("7. Demais telas do protótipo (implementadas)", level=1)
para(doc, "As cinco telas do protótipo de 16/10 já funcionam no PoC, além de Aulas (dados extraídos pelo scraper).", color=MUTED)
shots_row(
    doc,
    [("07-disciplinas.png", "Tela 3 — Disciplinas"), ("09-progresso.png", "Tela 5 — Progresso"), ("08-disciplina-detalhe.png", "Detalhe da disciplina")],
    width_cm=3.9,
)
shots_row(doc, [("14-tutor.png", "Tela 4 — Professor Tutor IA (desktop)"), ("13-aulas.png", "Aulas sincronizadas (desktop)")], width_cm=7.0)

doc.add_heading("8. Dicas de tela da aula aplicadas", level=1)
table(
    doc,
    ["Dica", "Como foi aplicada"],
    [
        ["Telas com muitos recursos confundem no celular", "Uma tarefa por tela e no máximo 5 destinos na navegação inferior."],
        ["Interface minimalista com chamadas para outras telas", "Cartões que levam ao detalhe; ações secundárias em menus e modais."],
        ["Telas que precisam de ajuda/treinamento", "Apresentação inicial de 3 passos e botão “?” em todas as telas."],
        ["Símbolos já consagrados", "Ícones padrão: lixeira para excluir, lápis para editar, olho para mostrar a senha, seta para enviar."],
        ["Informação rápida e completa", "A tela Início reúne as metas do dia, o progresso por disciplina e as próximas aulas."],
    ],
    widths=[6.5, 10.5],
)

# O parágrafo exigido após a última tabela não pode abrir uma página em branco.
tail = doc.paragraphs[-1]
tail.paragraph_format.space_before = Pt(0)
tail.paragraph_format.space_after = Pt(0)
tail.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
tail.paragraph_format.line_spacing = Pt(1)

doc.save(OUT)
print(f"Gerado: {OUT}")
